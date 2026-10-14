from odoo import models, fields, api
import base64
import xlrd
import io
import csv
from openpyxl import load_workbook
from odoo.exceptions import UserError, ValidationError
from datetime import datetime
import pytz


class SnkStagingSaleWizard(models.TransientModel):
    _name = "snk.staging.sale.wizard"
    _description = "Staging Sale"

    file = fields.Binary(string="File", required=True)
    filename = fields.Char("Filename")

    def action_confirm(self):
        self.ensure_one()
        if not self.file:
            return

        try:
            data = base64.b64decode(self.file)
            workbook = load_workbook(io.BytesIO(data), data_only=True)
            sheet = workbook.active

            StagingOrder = self.env['snk.staging.sale']
            StagingOrderLine = self.env['snk.staging.sale.line']

            for row in sheet.iter_rows(min_row=2, values_only=True):
                order_reference = self.safe_get(row, 0)
                order_date = self.safe_get(row, 1)
                customer = self.safe_get(row, 2, False)
                brand = self.safe_get(row, 3, False)
                warehouse = self.safe_get(row, 4, False)
                product = self.safe_get(row, 5, False)
                qty = self.safe_get(row, 6, 0)
                unit_price = self.safe_get(row, 7, 0)
                discount_amount = self.safe_get(row, 8, 0) or 0
                has_voucher = self.safe_get(row, 9, False)
                delivery_method = self.safe_get(row, 10, False)
                city = self.safe_get(row, 11, False)
                shipping_weight = self.safe_get(row, 12, False)
                snk_site_id = self.safe_get(row, 13, False)

                status = 'pending'

                if order_date:
                    tz = pytz.timezone('Asia/Jakarta')
                    if isinstance(order_date, str):
                        try:
                            order_date = datetime.strptime(order_date, "%Y-%m-%d %H:%M:%S")
                        except ValueError:
                            try:
                                order_date = datetime.strptime(order_date, "%Y-%m-%d %H:%M")
                            except ValueError:
                                order_date = datetime.strptime(order_date, "%d %b %Y %H:%M")

                    order_date = tz.localize(order_date).astimezone(pytz.UTC)
                    order_date = order_date.replace(tzinfo=None)

                staging = StagingOrder.search([('name', '=', order_reference)], limit=1)
                if not staging and order_reference:
                    staging = StagingOrder.create({
                        'name': order_reference,
                        'date_order': order_date,
                        'partner_id': customer,
                        'analytic_account_ids': brand,
                        'warehouse_id': warehouse,
                        'carrier_id': delivery_method,
                        'city_id': city,
                        'shipping_weight': shipping_weight or 0,
                        'snk_site_id': snk_site_id,
                        'status': status,
                    })

                line = StagingOrderLine.search([
                    ('order_reference', '=', staging.name),
                    ('product_code', '=', product),
                ], limit=1)
                if not line and staging:
                    StagingOrderLine.create({
                        'staging_sale_id': staging.id,
                        'order_reference': staging.name,
                        'product_code': product,
                        'analytic_id': brand,
                        'warehouse_id': warehouse,
                        'quantity': qty or 0,
                        'price_unit': unit_price or 0,
                        'snk_has_voucher': has_voucher,
                        'snk_discount_amount': discount_amount,
                        'status': status,
                    })
        except Exception as e:
            raise ValidationError(str(e))
    def _parse_weight(self, weight):
        if not weight:
            return 0.0

        w = str(weight).strip().lower()

        w = w.replace(",", ".")

        if "gr" in w:
            try:
                return float(w.replace("gr", "").strip()) / 1000.0
            except:
                return 0.0

        if "kg" in w:
            try:
                return float(w.replace("kg", "").strip())
            except:
                return 0.0

        try:
            return float(w)
        except:
            return 0.0

    def safe_get(self, row, idx, default=None):
        return row[idx] if idx < len(row) else default