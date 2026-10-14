from odoo import models, fields, api
import base64
import xlrd
import io
from openpyxl import load_workbook
from odoo.exceptions import UserError, ValidationError


class SnkSaleOrderFilterLine(models.TransientModel):
    _name = 'snk.sale.order.filter.line'
    _description = "Order Number Lines"

    wizard_id = fields.Many2one('snk.sale.order.filter', ondelete='cascade')
    name = fields.Char("Number")


class SnkSaleOrderFilter(models.TransientModel):
    _name = 'snk.sale.order.filter'
    _description = "To Order by ID"

    upload_file = fields.Binary("Upload Excel")
    file_name = fields.Char("File Name")
    line_ids = fields.One2many('snk.sale.order.filter.line', 'wizard_id', string="Numbers")

    @api.onchange('upload_file')
    def _onchange_upload_file(self):
        if not self.upload_file:
            self.line_ids = [(5, 0, 0)]
            return

        data = base64.b64decode(self.upload_file)
        wb = load_workbook(filename=io.BytesIO(data), read_only=True, data_only=True)
        sheet = wb.active

        header_keywords = ('order reference', 'number', 'no', 'order', 'order number', 'order_id', 'order id')

        line_vals = [(5, 0, 0)]
        col_index = None
        first_row = True

        for row in sheet.iter_rows(values_only=True):
            if not row:
                continue

            if first_row:
                for idx, cell_value in enumerate(row):
                    if cell_value and any(k in str(cell_value).lower() for k in header_keywords):
                        col_index = idx
                        break
                first_row = False
                continue

            if col_index is not None and col_index < len(row):
                value = str(row[col_index]).strip() if row[col_index] else ''
                if value:
                    line_vals.append((0, 0, {'name': value}))

        self.line_ids = line_vals

    def action_confirm(self):
        numbers = []
        for line in self.line_ids:
            if line.name:
                numbers.append(line.name)

        if not numbers:
            raise UserError("No order numbers found!")

        action = self.env.ref('sale.action_orders').read()[0]
        # action = self.env.ref('sale.action_quotations_with_onboarding').read()[0]

        action['domain'] = [('name', 'in', numbers)]

        return action